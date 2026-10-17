
from pathlib import Path

import pandas as pd
from app.rag.ocr import extract_text_from_image
from docx import Document
from pypdf import PdfReader
from pdf2image import convert_from_path
import pytesseract


SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".csv", ".xls", ".xlsx",
    ".jpg", ".jpeg", ".png",
}


def extract_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)
    pages = []

    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text.strip())

    # OCR fallback for scanned PDFs
    if not any(pages):
        images = convert_from_path(file_path)
        for image in images:
            text = pytesseract.image_to_string(image)
            if text.strip():
                pages.append(text.strip())

    return "\n\n".join(pages)


def extract_docx(file_path: str) -> str:
    document = Document(file_path)
    paragraphs = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraphs.append(text)

    for table in document.tables:
        for row in table.rows:
            row_text = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]
            if row_text:
                paragraphs.append(" | ".join(row_text))

    return "\n\n".join(paragraphs)


def extract_txt(file_path: str) -> str:
    return Path(file_path).read_text(
        encoding="utf-8",
        errors="ignore",
    )


def extract_csv(file_path: str) -> str:
    return pd.read_csv(file_path).to_csv(index=False)


def extract_excel(file_path: str) -> str:
    excel_file = pd.ExcelFile(file_path)
    sheets = []

    for sheet_name in excel_file.sheet_names:
        dataframe = pd.read_excel(
            file_path,
            sheet_name=sheet_name,
        )
        sheets.append(
            f"Sheet: {sheet_name}\n{dataframe.to_csv(index=False)}"
        )

    return "\n\n".join(sheets)


def extract_text_from_file(file_path: str) -> str:
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {extension}")

    if extension == ".pdf":
        return extract_pdf(file_path)

    if extension == ".docx":
        return extract_docx(file_path)

    if extension == ".txt":
        return extract_txt(file_path)

    if extension == ".csv":
        return extract_csv(file_path)

    if extension in {".xls", ".xlsx"}:
        return extract_excel(file_path)

    if extension in {".jpg", ".jpeg", ".png"}:
        return extract_text_from_image(file_path)

    raise ValueError(f"Unsupported file type: {extension}")
