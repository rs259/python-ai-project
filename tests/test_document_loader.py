from app.rag.document_loader import extract_text_from_file
from docx import Document
import csv


def test_extract_text_from_txt(tmp_path):
    file = tmp_path / "sample.txt"
    file.write_text(
        "India currency is Indian Rupee. Code INR.",
        encoding="utf-8",
    )

    result = extract_text_from_file(str(file))

    assert "Indian Rupee" in result
    assert "INR" in result


def test_extract_text_from_pdf(tmp_path):
    from reportlab.pdfgen import canvas

    file = tmp_path / "sample.pdf"
    pdf = canvas.Canvas(str(file))
    pdf.drawString(100, 750, "India currency is Indian Rupee INR")
    pdf.save()

    result = extract_text_from_file(str(file))

    assert "Indian Rupee" in result
    assert "INR" in result


def test_extract_text_from_docx(tmp_path):
    file = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("India currency is Indian Rupee INR")
    document.save(file)

    result = extract_text_from_file(str(file))

    assert "Indian Rupee" in result
    assert "INR" in result


def test_extract_text_from_csv(tmp_path):
    file = tmp_path / "sample.csv"

    with open(file, "w", newline="", encoding="utf-8") as csvfile:
        writer = csv.writer(csvfile)
        writer.writerow(["Country", "Currency", "Code"])
        writer.writerow(["India", "Indian Rupee", "INR"])

    result = extract_text_from_file(str(file))

    assert "Indian Rupee" in result
    assert "INR" in result


def test_extract_text_from_excel(tmp_path):
    from openpyxl import Workbook

    file = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Country", "Currency", "Code"])
    sheet.append(["India", "Indian Rupee", "INR"])
    workbook.save(file)

    result = extract_text_from_file(str(file))

    assert "Indian Rupee" in result
    assert "INR" in result


def test_extract_text_from_image(tmp_path):
    from PIL import Image, ImageDraw, ImageFont

    file = tmp_path / "sample.png"
    image = Image.new("RGB", (1000, 220), "white")
    draw = ImageDraw.Draw(image)

    try:
        font = ImageFont.truetype(
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            48,
        )
    except OSError:
        font = ImageFont.load_default()

    draw.text(
        (30, 60),
        "India currency is INR",
        fill="black",
        font=font,
    )
    image.save(file)

    result = extract_text_from_file(str(file)).upper()

    assert "INDIA" in result
    assert "INR" in result



def test_extract_text_from_scanned_pdf(tmp_path):
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader

    image_path = tmp_path / "scanned_page.png"
    pdf_path = tmp_path / "scanned.pdf"

    image = Image.new("RGB", (1200, 250), "white")
    draw = ImageDraw.Draw(image)

    try:
        font = ImageFont.truetype(
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            60,
        )
    except OSError:
        font = ImageFont.load_default()

    draw.text(
        (40, 70),
        "SCANNED TEST INDIA INR",
        font=font,
        fill="black",
    )
    image.save(image_path)

    pdf = canvas.Canvas(str(pdf_path))
    pdf.drawImage(
        ImageReader(image),
        0, 0,
        width=600,
        height=125,
    )
    pdf.save()

    result = extract_text_from_file(str(pdf_path)).upper()

    assert "SCANNED" in result
    assert "INDIA" in result
    assert "INR" in result
